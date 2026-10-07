import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


class ResumeExtractor:

    # =========================================================
    # PDF: LAYOUT-AWARE EXTRACTION
    # =========================================================

    @classmethod
    def extract_pdf_blocks(
        cls,
        file_path: str
    ) -> List[Dict[str, Any]]:
        """
        Extract PDF text while preserving basic layout information.
        """

        import fitz

        document: List[Dict[str, Any]] = []

        pdf = fitz.open(file_path)

        try:
            for page_number, page in enumerate(pdf, start=1):
                page_dict = page.get_text("dict")

                for block in page_dict.get("blocks", []):

                    # Ignore images and non-text blocks.
                    if block.get("type") != 0:
                        continue

                    for line in block.get("lines", []):
                        spans = line.get("spans", [])

                        if not spans:
                            continue

                        line_text = "".join(
                            str(span.get("text", ""))
                            for span in spans
                        ).strip()

                        if not line_text:
                            continue

                        font_sizes = [
                            float(span.get("size", 0))
                            for span in spans
                            if span.get("size") is not None
                        ]

                        max_font_size = (
                            max(font_sizes)
                            if font_sizes
                            else 0.0
                        )

                        is_bold = any(
                            "bold" in str(
                                span.get("font", "")
                            ).lower()
                            for span in spans
                        )

                        bbox = line.get(
                            "bbox",
                            block.get(
                                "bbox",
                                [0, 0, 0, 0]
                            )
                        )

                        document.append({
                            "page": page_number,
                            "text": line_text,
                            "font_size": round(
                                max_font_size,
                                2
                            ),
                            "bold": is_bold,
                            "x0": round(
                                float(bbox[0]),
                                2
                            ),
                            "y0": round(
                                float(bbox[1]),
                                2
                            ),
                            "x1": round(
                                float(bbox[2]),
                                2
                            ),
                            "y1": round(
                                float(bbox[3]),
                                2
                            ),
                            "type": "text",
                        })

        finally:
            pdf.close()

        return document

    # =========================================================
    # PDF: CLASSIFY BLOCKS
    # =========================================================

    @classmethod
    def classify_pdf_blocks(
        cls,
        blocks: List[Dict[str, Any]]
    ) -> List[Dict[str, Any]]:

        if not blocks:
            return []

        font_sizes = [
            float(block.get("font_size", 0))
            for block in blocks
            if float(block.get("font_size", 0)) > 0
        ]

        if font_sizes:
            sorted_sizes = sorted(font_sizes)
            body_font_size = sorted_sizes[
                len(sorted_sizes) // 2
            ]
        else:
            body_font_size = 10.0

        classified: List[Dict[str, Any]] = []

        for block in blocks:

            text = str(
                block.get("text", "")
            ).strip()

            if not text:
                continue

            font_size = float(
                block.get(
                    "font_size",
                    body_font_size
                )
            )

            bold = bool(
                block.get("bold", False)
            )

            # -------------------------------------------------
            # Bullet detection
            # -------------------------------------------------

            is_bullet = bool(
                re.match(
                    r"^(?:[•●▪◦‣⁃\-*]|\d+[.)])\s+",
                    text
                )
            )

            # -------------------------------------------------
            # Heading detection
            # -------------------------------------------------

            words = text.split()

            short_line = (
                len(words) <= 12
                and len(text) <= 100
            )

            larger_font = (
                font_size >= body_font_size * 1.20
            )

            bold_line = (
                bold
                and font_size >= body_font_size * 1.05
            )

            looks_like_heading = (
                short_line
                and (
                    larger_font
                    or bold_line
                )
            )

            if is_bullet:
                block_type = "bullet"
            elif looks_like_heading:
                block_type = "heading"
            else:
                block_type = "text"

            classified_block = dict(block)
            classified_block["type"] = block_type

            classified.append(
                classified_block
            )

        return classified

    # =========================================================
    # PDF: BLOCKS -> MARKDOWN
    # =========================================================

    @classmethod
    def blocks_to_markdown(
        cls,
        blocks: List[Dict[str, Any]]
    ) -> str:

        lines: List[str] = []

        for block in blocks:

            text = str(
                block.get("text", "")
            ).strip()

            if not text:
                continue

            block_type = block.get(
                "type",
                "text"
            )

            if block_type == "heading":

                lines.append(
                    f"## {text}"
                )

            elif block_type == "bullet":

                cleaned = re.sub(
                    r"^(?:[•●▪◦‣⁃\-*]|\d+[.)])\s+",
                    "",
                    text
                )

                lines.append(
                    f"- {cleaned}"
                )

            else:

                lines.append(text)

        return "\n\n".join(lines)

    # =========================================================
    # PDF: COMPLETE STRUCTURED EXTRACTION
    # =========================================================

    @classmethod
    def extract_pdf_document(
        cls,
        file_path: str
    ) -> Dict[str, Any]:

        try:

            blocks = cls.extract_pdf_blocks(
                file_path
            )

            classified_blocks = (
                cls.classify_pdf_blocks(
                    blocks
                )
            )

            raw_text = "\n".join(
                block["text"]
                for block in classified_blocks
                if block.get("text")
            )

            markdown_text = (
                cls.blocks_to_markdown(
                    classified_blocks
                )
            )

            word_count = len(
                raw_text.split()
            )

            # -------------------------------------------------
            # OCR fallback
            # -------------------------------------------------

            if word_count < 50:

                ocr_text, ocr_confidence = (
                    cls._run_ocr(file_path)
                )

                if (
                    ocr_text
                    and len(ocr_text.split())
                    > word_count
                ):

                    return {
                        "raw_text": ocr_text,
                        "markdown_text": ocr_text,
                        "blocks": [],
                        "ocr_applied": True,
                        "ocr_confidence": (
                            ocr_confidence
                        ),
                        "ocr_error": None,
                    }

            return {
                "raw_text": raw_text,
                "markdown_text": markdown_text,
                "blocks": classified_blocks,
                "ocr_applied": False,
                "ocr_confidence": None,
                "ocr_error": None,
            }

        except Exception as e:

            # IMPORTANT:
            # Do not silently hide the extraction error.
            print(
                "PDF EXTRACTION ERROR:",
                repr(e)
            )

            return {
                "raw_text": "",
                "markdown_text": "",
                "blocks": [],
                "ocr_applied": False,
                "ocr_confidence": None,
                "ocr_error": str(e),
            }

    # =========================================================
    # PUBLIC FILE EXTRACTION
    # =========================================================

    @staticmethod
    def extract_from_file(
        file_path: str,
        filename: str
    ) -> Tuple[
        str,
        bool,
        Optional[float],
        Optional[str]
    ]:

        ext = Path(
            filename
        ).suffix.lower()

        ocr_applied = False
        ocr_confidence = None
        error = None
        text = ""

        try:

            # -------------------------------------------------
            # PDF
            # -------------------------------------------------

            if ext == ".pdf":

                result = (
                    ResumeExtractor
                    .extract_pdf_document(
                        file_path
                    )
                )

                text = result["raw_text"]

                ocr_applied = (
                    result["ocr_applied"]
                )

                ocr_confidence = (
                    result["ocr_confidence"]
                )

                error = (
                    result["ocr_error"]
                )

            # -------------------------------------------------
            # DOCX / DOC
            # -------------------------------------------------

            elif ext in [".docx", ".doc"]:

                text = (
                    ResumeExtractor
                    ._extract_docx(
                        file_path
                    )
                )

            # -------------------------------------------------
            # TXT / Markdown
            # -------------------------------------------------

            elif ext in [
                ".txt",
                ".md",
                ".markdown"
            ]:

                with open(
                    file_path,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as f:

                    text = f.read()

            # -------------------------------------------------
            # Images
            # -------------------------------------------------

            elif ext in [
                ".png",
                ".jpg",
                ".jpeg",
                ".tiff",
                ".bmp"
            ]:

                text, ocr_confidence = (
                    ResumeExtractor
                    ._run_ocr(
                        file_path
                    )
                )

                ocr_applied = True

            # -------------------------------------------------
            # Unknown file
            # -------------------------------------------------

            else:

                with open(
                    file_path,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as f:

                    text = f.read()

        except Exception as e:

            error = str(e)

        # -----------------------------------------------------
        # Normalize extracted text
        # -----------------------------------------------------

        if text:

            clean_text = (
                text
                .replace("\r\n", "\n")
                .replace("\r", "\n")
            )

            clean_text = re.sub(
                r"[ \t]+",
                " ",
                clean_text
            )

            clean_text = re.sub(
                r"\n{3,}",
                "\n\n",
                clean_text
            )

            clean_text = "\n".join(
                line.strip()
                for line in clean_text.split("\n")
            ).strip()

        else:

            clean_text = ""

        return (
            clean_text,
            ocr_applied,
            ocr_confidence,
            error
        )

    # =========================================================
    # LEGACY PDF EXTRACTION
    # =========================================================

    @staticmethod
    def _extract_pdf(
        file_path: str
    ) -> Tuple[
        str,
        bool,
        Optional[float]
    ]:

        result = (
            ResumeExtractor
            .extract_pdf_document(
                file_path
            )
        )

        return (
            result["raw_text"],
            result["ocr_applied"],
            result["ocr_confidence"],
        )

    # =========================================================
    # DOCX
    # =========================================================

    @staticmethod
    def _extract_docx(
        file_path: str
    ) -> str:

        import docx

        doc = docx.Document(
            file_path
        )

        paragraphs = [
            p.text
            for p in doc.paragraphs
            if p.text.strip()
        ]

        for table in doc.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():

                        paragraphs.append(
                            cell.text.strip()
                        )

        return "\n".join(
            paragraphs
        )

    # =========================================================
    # OCR
    # =========================================================

    @staticmethod
    def _run_ocr(
        file_path: str
    ) -> Tuple[
        str,
        Optional[float]
    ]:

        try:

            import pytesseract
            from PIL import Image

            ext = Path(
                file_path
            ).suffix.lower()

            # -------------------------------------------------
            # Image OCR
            # -------------------------------------------------

            if ext in [
                ".png",
                ".jpg",
                ".jpeg",
                ".tiff",
                ".bmp"
            ]:

                img = Image.open(
                    file_path
                )

                text = (
                    pytesseract
                    .image_to_string(img)
                )

                return text, 0.85

            # -------------------------------------------------
            # PDF OCR
            # -------------------------------------------------

            elif ext == ".pdf":

                # PDF OCR will be implemented later.
                return (
                    "",
                    None
                )

        except Exception as e:

            print(
                "OCR ERROR:",
                repr(e)
            )

        return "", None